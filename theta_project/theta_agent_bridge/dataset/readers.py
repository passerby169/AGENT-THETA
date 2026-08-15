from __future__ import annotations

import csv
import codecs
import hashlib
import json
import random
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Iterator


SUPPORTED_SUFFIXES = {'.csv', '.tsv', '.txt', '.json', '.jsonl', '.xlsx', '.xls', '.parquet'}
DEFAULT_PROFILE_LIMIT = 5000
READER_VERSION = '2.0.0'


@dataclass
class DatasetReader:
    path: Path
    suffix: str
    columns: list[str]
    encoding: str
    delimiter: str | None
    row_count: int
    rows: list[dict[str, Any]]
    head_rows: list[dict[str, Any]]
    rows_truncated: bool
    sheets: list[str]
    selected_sheet: str | None
    warnings: list[str]


def load_dataset(
    path: Path,
    *,
    seed: str = 'theta-dataset-profile',
    profile_limit: int = DEFAULT_PROFILE_LIMIT,
    sheet_name: str | None = None,
) -> DatasetReader:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported dataset suffix '{suffix}'. Supported: {sorted(SUPPORTED_SUFFIXES)}")
    if suffix == '.xlsx':
        return _load_xlsx_streaming(path, seed, profile_limit, sheet_name)
    if suffix in {'.xls', '.parquet'}:
        return _load_dataframe(path, suffix, seed, profile_limit, sheet_name)
    if suffix in {'.csv', '.tsv'}:
        encoding = _detect_encoding(path)
        delimiter = '\t' if suffix == '.tsv' else _sniff_delimiter(path, encoding)
        with path.open('r', encoding=encoding, newline='', errors='strict') as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            columns = _unique_column_names(next(reader, []))
            records = (
                {column: row[index] if index < len(row) else None for index, column in enumerate(columns)}
                for row in reader
            )
            result = _bounded_reader(path, suffix, columns, encoding, delimiter, records, seed, profile_limit)
            _append_encoding_warning(result)
            return result
    if suffix == '.jsonl':
        encoding = _detect_encoding(path)

        def records() -> Iterator[dict[str, Any]]:
            with path.open('r', encoding=encoding, errors='strict') as handle:
                for line in handle:
                    if not line.strip():
                        continue
                    value = json.loads(line)
                    if not isinstance(value, dict):
                        raise ValueError('JSONL records must be objects')
                    yield value

        return _bounded_reader(path, suffix, None, encoding, None, records(), seed, profile_limit)
    if suffix == '.json':
        encoding = _detect_encoding(path)
        with path.open('r', encoding=encoding, errors='strict') as handle:
            value = json.load(handle)
        records = value.get('data') or value.get('records') or value.get('items') if isinstance(value, dict) else value
        if not isinstance(records, list):
            raise ValueError('JSON dataset must be a list or contain data/records/items')
        object_rows = (item for item in records if isinstance(item, dict))
        return _bounded_reader(path, suffix, None, encoding, None, object_rows, seed, profile_limit)

    encoding = _detect_encoding(path)

    def text_records() -> Iterator[dict[str, Any]]:
        with path.open('r', encoding=encoding, errors='strict') as handle:
            for line in handle:
                value = line.strip()
                if value:
                    yield {'text': value}

    return _bounded_reader(path, suffix, ['text'], encoding, None, text_records(), seed, profile_limit)


def _bounded_reader(
    path: Path,
    suffix: str,
    known_columns: list[str] | None,
    encoding: str,
    delimiter: str | None,
    records: Iterable[dict[str, Any]],
    seed: str,
    profile_limit: int,
) -> DatasetReader:
    limit = max(10, profile_limit)
    rng = random.Random(_seed_value(seed))
    columns = list(known_columns or [])
    seen_columns = set(columns)
    rows: list[dict[str, Any]] = []
    head_rows: list[dict[str, Any]] = []
    row_count = 0
    for raw_row in records:
        row = {_normalize_column_name(key): _json_safe(value) for key, value in raw_row.items()}
        for key in row:
            if key not in seen_columns:
                columns.append(key)
                seen_columns.add(key)
        if len(head_rows) < 10:
            head_rows.append(row)
        row_count += 1
        if len(rows) < limit:
            rows.append(row)
            continue
        replacement = rng.randrange(row_count)
        if replacement < limit:
            rows[replacement] = row
    return DatasetReader(
        path,
        suffix,
        columns,
        encoding,
        delimiter,
        row_count,
        rows,
        head_rows,
        row_count > len(rows),
        [],
        None,
        [],
    )


def _load_xlsx_streaming(
    path: Path,
    seed: str,
    profile_limit: int,
    sheet_name: str | None,
) -> DatasetReader:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError('Reading .xlsx requires the optional openpyxl adapter.') from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [str(name) for name in workbook.sheetnames]
        if not sheets:
            raise ValueError('XLSX workbook does not contain any worksheets')
        if sheet_name and sheet_name not in sheets:
            raise ValueError(f"Excel sheet '{sheet_name}' was not found. Available sheets: {sheets}")
        candidates = [sheet_name] if sheet_name else sheets
        selected_sheet = candidates[0]
        header: tuple[Any, ...] = ()
        records_iter: Iterator[tuple[Any, ...]] = iter(())
        for candidate in candidates:
            worksheet = workbook[candidate]
            iterator = worksheet.iter_rows(values_only=True)
            first = next(iterator, ())
            selected_sheet = candidate
            header = first
            records_iter = iterator
            if any(value is not None and str(value).strip() for value in first):
                break
        columns = _unique_column_names(header)

        def records() -> Iterator[dict[str, Any]]:
            for values in records_iter:
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                yield {
                    column: _json_safe(values[index] if index < len(values) else None)
                    for index, column in enumerate(columns)
                }

        reader = _bounded_reader(
            path,
            '.xlsx',
            columns,
            'binary',
            None,
            records(),
            seed,
            profile_limit,
        )
        reader.sheets = sheets
        reader.selected_sheet = selected_sheet
        return reader
    finally:
        workbook.close()


def _load_dataframe(
    path: Path,
    suffix: str,
    seed: str,
    profile_limit: int,
    sheet_name: str | None,
) -> DatasetReader:
    try:
        import pandas as pd
    except ImportError as exc:
        raise RuntimeError(
            f"Reading {suffix} requires the optional pandas adapter and its format engine."
        ) from exc
    sheets: list[str] = []
    selected_sheet: str | None = None
    if suffix == '.parquet':
        frame = pd.read_parquet(path)
    else:
        workbook = pd.ExcelFile(path)
        sheets = [str(name) for name in workbook.sheet_names]
        if sheet_name:
            if sheet_name not in sheets:
                raise ValueError(f"Excel sheet '{sheet_name}' was not found. Available sheets: {sheets}")
            selected_sheet = sheet_name
            frame = pd.read_excel(workbook, sheet_name=sheet_name)
        else:
            frame = None
            for candidate in sheets:
                candidate_frame = pd.read_excel(workbook, sheet_name=candidate)
                if not candidate_frame.empty:
                    selected_sheet = candidate
                    frame = candidate_frame
                    break
            if frame is None:
                selected_sheet = sheets[0] if sheets else None
                frame = pd.DataFrame()
    frame.columns = _unique_column_names(frame.columns)
    frame = frame.where(frame.notna(), None)
    records = frame.to_dict(orient='records')
    reader = _bounded_reader(
        path,
        suffix,
        [str(column) for column in frame.columns],
        'binary',
        None,
        records,
        seed,
        profile_limit,
    )
    reader.sheets = sheets
    reader.selected_sheet = selected_sheet
    return reader


def _detect_encoding(path: Path) -> str:
    with path.open('rb') as handle:
        prefix = handle.read(4)
        handle.seek(0)
        sample = handle.read(64 * 1024)
    if prefix.startswith(codecs.BOM_UTF8):
        return 'utf-8-sig'
    if prefix.startswith(codecs.BOM_UTF16_LE):
        return 'utf-16-le'
    if prefix.startswith(codecs.BOM_UTF16_BE):
        return 'utf-16-be'
    for encoding in ('utf-8', 'gb18030', 'gbk'):
        try:
            codecs.getincrementaldecoder(encoding)(errors='strict').decode(sample, final=False)
            return encoding
        except UnicodeDecodeError:
            continue
    return 'latin-1'


def _append_encoding_warning(reader: DatasetReader) -> None:
    if reader.encoding == 'latin-1':
        reader.warnings.append(
            '文本编码无法可靠识别，已使用 latin-1 保底读取；请复核中文列名和样本内容。'
        )


def _sniff_delimiter(path: Path, encoding: str) -> str:
    with path.open('r', encoding=encoding, errors='strict') as handle:
        sample = ''.join(handle.readline() for _ in range(20))
    try:
        return csv.Sniffer().sniff(sample, delimiters=',\t;|').delimiter
    except csv.Error:
        return ','


def _seed_value(seed: str) -> int:
    return int(hashlib.sha256(seed.encode('utf-8')).hexdigest()[:16], 16)


def _normalize_column_name(value: Any) -> str:
    text = unicodedata.normalize('NFC', str(value if value is not None else ''))
    text = text.lstrip('\ufeff').replace('\u200b', '').replace('\u200c', '').replace('\u200d', '')
    text = ''.join(character for character in text if unicodedata.category(character) != 'Cc')
    return text.strip()


def _unique_column_names(values: Iterable[Any]) -> list[str]:
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values):
        base = _normalize_column_name(value) or f'column_{index + 1}'
        seen[base] = seen.get(base, 0) + 1
        names.append(base if seen[base] == 1 else f'{base} [{seen[base]}]')
    return names


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)
